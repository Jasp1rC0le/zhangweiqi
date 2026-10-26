"""Read public Office originals without macros, external links or publication.

Review flags are intentionally untouched. ZIP + supplied extracted folder are one source.
"""
from pathlib import Path
import hashlib
import json
import re
import zipfile
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RECEIPT = ROOT/'docs/research/additional-attachments-received-2026-10-11.json'
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def mask_phone(text):
    return re.sub(r'(?<!\d)1[3-9]\d{9}(?!\d)', '[私人手机号不公开]', str(text))


def inspect_office(path):
    details = {'extracted': True}
    if path.suffix == '.docx':
        document = Document(path)
        blocks = []
        for block in document.element.body:
            if block.tag == W+'p':
                blocks.append(''.join(node.text or '' for node in block.iter(W+'t')))
            elif block.tag == W+'tbl':
                for row in block.iter(W+'tr'):
                    blocks.append(' | '.join(''.join(node.text or '' for node in cell.iter(W+'t')) for cell in row.findall(W+'tc')))
        text = '\n'.join(blocks)
        images = []
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist():
                if not info.filename.startswith('word/media/') or info.file_size > 10*1024*1024:
                    continue
                suffix = Path(info.filename).suffix.lower()
                if suffix not in ('.png','.jpg','.jpeg','.gif','.emf','.wmf'):
                    continue
                target = path.with_name(path.stem+'-'+Path(info.filename).name)
                payload = archive.read(info)
                if target.exists() and hashlib.sha256(target.read_bytes()).digest() != hashlib.sha256(payload).digest():
                    raise ValueError('Preserve changed embedded original')
                if not target.exists(): target.write_bytes(payload)
                images.append(str(target.relative_to(ROOT)))
        details.update(embeddedImages=images, paragraphs=len(document.paragraphs), tables=len(document.tables),
                       renderVerified=False, note='按正文顺序提取全部段落和表格，内嵌原图需另行核读；不执行外部关系或宏')
    elif path.suffix == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        blocks, sheets = [], []
        try:
            for sheet in book:
                if sheet.max_row > 500 or sheet.max_column > 60: raise ValueError('Unexpected sheet bounds')
                rows = []
                for row in sheet:
                    cells = [f'{cell.coordinate}={mask_phone(cell.value)}' for cell in row if cell.value is not None]
                    if cells: rows.append(' | '.join(cells))
                blocks.append(f'=== {sheet.title} state={sheet.sheet_state} ===\n'+'\n'.join(rows))
                sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                               'nonemptyRows': len(rows), 'state': sheet.sheet_state})
        finally: book.close()
        text = '\n'.join(blocks)
        details.update(sheets=sheets, note='全部有值单元格只读提取，私人手机号已遮蔽；未修改工作簿、未执行公式或外部链接')
    else: raise ValueError('Not an Office original')
    text = mask_phone(text)
    if not text.strip() or len(text)>200000: raise ValueError('Empty or unexpectedly large extraction')
    target = path.with_suffix('.txt')
    target.write_text(text)
    details.update(textArchive=str(target.relative_to(ROOT)), chars=len(text))
    return details


receipt = json.loads(RECEIPT.read_text())
for row in receipt['files']:
    path = ROOT/row['archive']
    if path.suffix in ('.xlsx','.docx'):
        row.update(inspect_office(path))
    elif path.suffix == '.zip':
        members = []
        directory = path.with_name('defence-2025-members')
        directory.mkdir(exist_ok=True)
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos)>50 or sum(info.file_size for info in infos)>100*1024*1024: raise ValueError('Unsafe ZIP size')
            for index, info in enumerate(infos):
                member = Path(info.filename)
                if member.is_absolute() or '..' in member.parts or (info.external_attr>>16)&0o170000 == 0o120000:
                    raise ValueError('Unsafe ZIP member')
                if info.is_dir(): continue
                if member.suffix.lower() not in ('.docx','.xlsx','.pdf'): raise ValueError('Unsupported ZIP member')
                payload = archive.read(info)
                digest = hashlib.sha256(payload).hexdigest()
                target = directory/f'member-{index+1}{member.suffix.lower()}'
                if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest()!=digest: raise ValueError('Preserve changed member')
                if not target.exists(): target.write_bytes(payload)
                prior = next((item for item in row.get('members',[]) if item.get('sha256')==digest),{})
                item = {**prior, 'name':info.filename,'sha256':digest,'archive':str(target.relative_to(ROOT))}
                if target.suffix in ('.xlsx','.docx'): item.update(inspect_office(target))
                members.append(item)
        row.update(members=members, extracted=all(item.get('extracted') for item in members),
                   note='ZIP内4个文件只读归档；提供的同一解压文件夹不重复入库，PDF与Word同正文合并审核')
RECEIPT.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps([{'key':row['key'],'extracted':row.get('extracted'),'chars':row.get('chars'),'images':len(row.get('embeddedImages',[])),'members':len(row.get('members',[]))} for row in receipt['files']],ensure_ascii=False))
