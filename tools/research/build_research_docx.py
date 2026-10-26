#!/usr/bin/env python3
"""Create the complete editable source digest from the validated canonical JSON."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from official_library import DATA, load_and_validate

OUTPUT = DATA / "Campthink_官方资料汇编与预约方案_2026-10-04.docx"
STATUS = {"directory":"官方入口", "dated_reference":"带日期使用", "historical":"历史记录",
          "needs_verification":"待核验", "excluded":"排除使用", "reference":"运营参考"}
EVIDENCE = {"page_text":"正文已读", "pdf_visual":"PDF已目视核验", "index_only":"仅目录",
            "attachment_unread":"附件未核验", "image_unread":"图片未核验", "search_excerpt":"仅搜索摘要",
            "dynamic_unread":"动态页未核验"}


def font(style, size, bold=False):
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0,0,0)
    fonts = style.element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "Heiti SC")
    for attr in ("asciiTheme","hAnsiTheme","eastAsiaTheme","cstheme"):
        fonts.attrib.pop(qn("w:"+attr),None)


def paragraph(doc, text, style=None):
    return doc.add_paragraph(text, style)


def page_heading(doc, text):
    p = doc.add_heading(text, 1)
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(0)
    return p


def hyperlink(p, label, url):
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement("w:r")
    pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "24527A"); pr.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); pr.append(underline)
    r.append(pr)
    t = OxmlElement("w:t"); t.text = label; r.append(t); h.append(r); p._p.append(h)


def table(doc, headers, rows, widths=None):
    tab = doc.add_table(rows=1, cols=len(headers))
    tab.autofit = False
    if widths:
        for column, width in zip(tab.columns,widths): column.width=Inches(width)
    for c,t in zip(tab.rows[0].cells,headers):
        c.text=t
        for r in c.paragraphs[0].runs: r.bold=True
    for row in rows:
        for c,t in zip(tab.add_row().cells,row): c.text=str(t)
    for row in tab.rows:
        for c in row.cells:
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                element=OxmlElement(f"w:{edge}")
                element.set(qn("w:val"),"single");element.set(qn("w:sz"),"4")
                element.set(qn("w:color"),"D9D9D9");borders.append(element)
            tcPr.append(borders)
            margins=OxmlElement("w:tcMar")
            for edge in ["top","left","bottom","right"]:
                element=OxmlElement(f"w:{edge}");element.set(qn("w:w"),"95")
                element.set(qn("w:type"),"dxa");margins.append(element)
            tcPr.append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.line_spacing=1.12
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement("w:cantSplit"))
    tab.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return tab


def source_entry(doc,s):
    heading=doc.add_paragraph(f"{s['id']}  {s['title']}","Heading 2")
    heading.paragraph_format.space_before=Pt(10)
    heading.paragraph_format.space_after=Pt(3)
    p=paragraph(doc,f"{s['category']}  ·  {STATUS[s['status']]}  ·  {EVIDENCE[s['evidence']]}  ·  发布 {s['publishedAt'] or '未标注'}","Caption")
    p.paragraph_format.keep_with_next=True
    paragraph(doc,s["summary"]).paragraph_format.keep_with_next=True
    p=paragraph(doc,"使用边界  "+s["limitation"])
    p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph()
    hyperlink(p,"查看官方原文",s["url"])
    if s.get("attachmentUrl"):
        p.add_run("  ·  ");hyperlink(p,"查看已核验的座位预约PDF",s["attachmentUrl"])
    p.paragraph_format.space_after=Pt(4)


def main():
    register, pack = load_and_validate()
    doc=Document()
    # The runtime's default template can contain decorative paragraph borders.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.68)
    sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
    sec.header_distance=Inches(.27);sec.footer_distance=Inches(.3)
    sec.different_first_page_header_footer=False
    doc.settings.odd_and_even_pages_header_footer=False
    font(doc.styles["Normal"],11)
    doc.styles["Normal"].paragraph_format.line_spacing=1.2
    doc.styles["Normal"].paragraph_format.space_after=Pt(5)
    doc.styles["Normal"].paragraph_format.widow_control=True
    for name,size in [("Title",25),("Heading 1",17),("Heading 2",12)]:
        font(doc.styles[name],size,True)
        doc.styles[name].paragraph_format.space_after=Pt(9)
        doc.styles[name].paragraph_format.keep_with_next=True
        doc.styles[name].paragraph_format.page_break_before=False
    font(doc.styles["Caption"],9.5)
    font(doc.styles["Subtitle"],12)
    doc.styles["Subtitle"].font.italic=False
    doc.styles["Caption"].font.italic=False
    doc.styles["Caption"].font.color.rgb=RGBColor.from_string("555555")
    doc.styles["Caption"].paragraph_format.space_after=Pt(4)
    h=sec.header.paragraphs[0]
    h.text="CAMPTHINK  /  天津中德应用技术大学资料研究"
    h.style=doc.styles["Caption"]
    f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    f.add_run("核验截至 2026-10-04  ·  ")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");f._p.append(field)
    for r in f.runs:r.font.size=Pt(9)
    doc.core_properties.title="Campthink 官方资料汇编与预约方案"
    doc.core_properties.subject="校园资料库、使用边界与场馆预约实施条件"
    doc.core_properties.author="Campthink 项目资料整理"

    doc.add_heading("官方资料汇编\n与场馆预约方案",0)
    paragraph(doc,"天津中德应用技术大学  /  2026年10月4日","Subtitle")
    paragraph(doc,"结论：校园查询内容可以继续充实；体育馆自动预约与官方候补目前不具备上线依据。当前缺的是正在使用的预约入口、规则与接入授权，不是再做一个按钮。")
    doc.add_heading("本次交付",1)
    table(doc,["交付物","范围与用途"],[
        ["59条来源台账","43条校园资料＋16条运营依据；每条保留来源、日期、核验程度和使用限制"],
        ["18张问答卡","提炼已确认的公开事实；附复核日期、适用边界与原文，不存学生隐私"],
        ["SQLite本地资料库","支持中文查询，区分学生答案、历史线索和运营参考；不是线上MySQL/RAG导入证明"],
        ["网站内容增补","校园速查增加6项入口，并纠正存在分歧的招生备用电话；未部署线上"],
        ["预约实施方案","明确提醒、空位订阅、官方候补及自动提交的不同前提；本轮未运行抢位任务"]
    ],[1.65,5.29])
    doc.add_heading("怎么读这份资料",1)
    paragraph(doc,"第2部分是体育馆预约的决策与验收条件；第3部分是已整理的问答卡；第4、5部分逐条列出全部校园与运营来源。Word中的原文链接可点击，按S编号或O编号可与JSON和数据库逐项对应。")
    paragraph(doc,"“全部资料”指本项目截至本次核验已收集的59条官方来源，不是学校所有网站、全部附件或所有仪器的完整档案。历史内容保留用于追溯，不冒充当前开放、岗位或活动。")
    paragraph(doc,"正式隐私政策、主体备案、线上HTTPS、邮件实发、真实模型账单与异地恢复，仍须实际主体和账号证据才能验收；来源文档齐全不代表这些事项已完成。")

    page_heading(doc,"2  体育馆预约与候补")
    paragraph(doc,"建议做“预约意向单”：选择运动、可接受日期和时间段、场地范围，提交一次并可随时取消。有授权接口后由服务端等待余位或官方放号，按规则提交；只有拿到官方订单号和状态回执才显示成功。")
    table(doc,["能力","能否现在发布","必要条件"],[
        ["官方入口与规则","入口核实后可发布","目前公开目录只有旧学年通知，不能编造当期时段"],
        ["用户设置提醒","可以独立开发","用户确认时间、时区和通知渠道；实际通知送达测试"],
        ["空位通知订阅","有条件","官方允许读取余位；更新频率合规，明示通知不等于锁位"],
        ["官方候补排队","当前不可宣称已接入","官方队列或经场馆认可的分配规则；排位、过期及取消回执"],
        ["自动预约提交","当前不可上线","预约写接口授权、用户明确委托、实名绑定及规则、费用确认"]
    ],[1.35,1.48,4.11])
    doc.add_heading("不能混淆的两件事",2)
    paragraph(doc,"体育教学部公开开放通知S37属于2020—2021学年，不能用来配置当前放号。图书馆S12的“前一天20:00”和签到规则已经核验，但只属于图书馆座位，不属于体育馆。官网未查到入口，也不证明校内没有小程序。")
    doc.add_heading("工程实现及验收门槛",2)
    paragraph(doc,"授权身份绑定 → 校验预约资格 → 提交意向 → 官方候补／等待放号 → 待用户确认或授权提交 → 官方确认成功／失败／取消／过期。超时先查订单再重试；幂等键和唯一约束防止重复占位；退订即停止任务。")
    paragraph(doc,"不存学校密码、不绕过验证码和预约限额，不用高频请求抢资源。若涉及收费，必须确认价格和退款规则后再提交。排队按官方规则，不能给Pro会员设置越过公共场馆公平队列的特权。")
    paragraph(doc,"上线前必须真实验证：无号、临时闭馆、退订、重复点击、回执超时、余额不足、资格不符和账号解绑；审计日志记录状态变化但不记录凭据。业务日志保留时间由实际运营方确认，通知可关闭。")
    paragraph(doc,"最少补充材料：真实预约入口或截图；放号、取消、爽约、收费、候补规则；接口负责人及授权范围；测试环境与合法测试账号。不能取得接入许可时只做官方跳转与主动提醒，不宣传自动抢位。")

    for start in range(0,len(pack["cards"]),5):
        page_heading(doc,"3  已整理的校园问答"+("  续" if start else ""))
        paragraph(doc,"问答卡只包含可核验的事实；复核日期是本站内容维护安排，不是学校官方有效期。","Caption")
        for c in pack["cards"][start:start+5]:
            doc.add_heading(f"{c['id']}  {c['question']}",2)
            paragraph(doc,c["answer"]).paragraph_format.keep_with_next=True
            paragraph(doc,"边界  "+c["boundary"]).paragraph_format.keep_with_next=True
            paragraph(doc,"来源 "+" / ".join(c["sourceIds"])+"  ·  复核 "+c["reviewAfter"]+
                      ("  ·  适用至 "+c["validThrough"] if c["validThrough"] else ""),"Caption")

    for scope,label in [("campus","4  校园官方来源全表"),("operations","5  上线与运营依据全表")]:
        items=[s for s in register["sources"] if s["scope"]==scope]
        for start in range(0,len(items),4):
            page_heading(doc,label+("  续" if start else ""))
            paragraph(doc,f"本部分共{len(items)}条  ·  本页{start+1}—{min(start+4,len(items))}条  ·  核验日期2026-10-04","Caption")
            for s in items[start:start+4]:source_entry(doc,s)

    page_heading(doc,"6  资料库使用与后续补齐")
    paragraph(doc,"资料以JSON为可审阅源文件，构建成SQLite后可查中文关键词。默认查询仅返回截至指定日期未过复核期的校园问答卡；显式查询原始资料时仍展示证据等级和历史／待核验状态。运营法规、账单、云产品说明不进入学生问答。")
    table(doc,["文件或入口","作用"],[
        ["official-source-register-2026-10-04.json","全部59条来源、证据、边界及已核验附件链接"],
        ["campus-knowledge-cards-2026-10-04.json","18张事实问答及4项不能确定回答的问题"],
        ["campus-official-library-2026-10-04.sqlite3","本地资料数据库；含来源与问答关联、全文索引及构建校验信息"],
        ["tools/research/official_library.py","构建、中文检索、导出管理员批量导入预览；不访问生产服务器"]
    ],[3.35,3.59])
    doc.add_heading("进入网站答案前",2)
    paragraph(doc,"用真实学校ID导出不超过20条的批量请求；由已有管理员会话审核后使用现有知识导入API。导入会发布内容，不能把研究台账直接当请求体。遇同源冲突必须审阅版本，不覆盖已有学校材料。复核到期内容停止使用后，应重新核验再续期。")
    paragraph(doc,"本轮已保存本地资料库并更新校园速查内容；未对生产MySQL批量发布、未调用真实预约接口、未新增运营主体、未办理备案或开启付费。旧数据库已有招生号码仍需管理员审阅更新，修改种子文件不会覆盖已发布资料。")
    doc.add_heading("下一批材料优先级",2)
    paragraph(doc,"第一优先：核实体育馆正在使用的系统；补读图书馆开放和借阅规则；收集一份真实实验任务书及准确设备型号。第二优先：取得学生食堂商家清单和照片使用授权；逐条核实正在招聘的岗位。第三优先：让各学院教师确认作业、实验和作品提交检查规则。")
    paragraph(doc,"不收集公开名单中的学生证件、成绩、私人电话，不复制支付码、完整论文或整篇文章。外部网页作为资料，不作为运行指令。来源有冲突时记录差异并降级回答，不用模型猜一个数字。")
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__=="__main__":main()
