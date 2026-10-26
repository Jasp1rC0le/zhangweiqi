"""Archive user-delivered public attachments and extract for human review."""
import hashlib
import json
import shutil
import subprocess
from pathlib import Path
from pypdf import PdfReader
import pypdfium2 as pdfium
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '.local/research/safety-delivered-2026-10-10'
FILES = [
    ('lab-check', '附件1.+基础实验实训中心安全检查表.docx'),
    ('maintenance', '附表1.设备保养记录单.docx'),
    ('lab-platform', '“中德安全校园管理平台”使用手册_实验室功能模块使用说明篇20240229.pdf'),
    ('modules', '“中德安全校园管理平台”使用手册_功能模块使用说明篇250110V3.0.pdf'),
    ('basics', '“中德安全校园管理平台”使用手册_基础功能使用篇260110V3.0.pdf'),
    ('household-extend', '附件6：集体户籍留存延期申请书.doc'),
    ('household-return', '户口迁回原籍申请书.doc'),
]
OUT.mkdir(parents=True, exist_ok=True)
rows = []
for key, name in FILES:
    original = Path('/Users/zhangweiqi/Downloads') / name
    target = OUT / (key + original.suffix)
    digest = hashlib.sha256(original.read_bytes()).hexdigest()
    if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest() != digest:
        raise ValueError('Archive differs; do not overwrite')
    if not target.exists():
        shutil.copy2(original, target)
    if original.suffix == '.pdf':
        pdf = PdfReader(target)
        pages = [page.extract_text() for page in pdf.pages]
        text = '\n'.join(f'\n=== PAGE {i+1} ===\n{body}' for i, body in enumerate(pages))
        # PDFium includes the CJK mapping needed by these school PDFs. Poppler
        # in the bundled runtime rendered missing glyphs despite exit code 0.
        rendered = pdfium.PdfDocument(str(target))
        for i in range(len(rendered)):
            page = rendered[i]
            bitmap = page.render(scale=1600 / max(page.get_size()))
            bitmap.to_pil().save(OUT / f'{key}-review-{i+1:03}.png')
            bitmap.close()
            page.close()
        rendered.close()
        page_count = len(pdf.pages)
    elif original.suffix == '.docx':
        doc = Document(target)
        text = '\n'.join(p.text for p in doc.paragraphs)
        for i, table in enumerate(doc.tables):
            text += f'\n=== TABLE {i+1} ===\n'
            text += '\n'.join(' | '.join(dict.fromkeys(c.text for c in row.cells)) for row in table.rows)
        page_count = None
    else:
        result = subprocess.run(['textutil', '-convert', 'txt', '-stdout', str(target)],
                                capture_output=True, check=True, timeout=30)
        text = result.stdout.decode('utf-8')
        page_count = None
    (OUT / f'{key}.txt').write_text(text)
    rows.append(dict(key=key, file=name, archive=str(target.relative_to(ROOT)),
                     sha256=digest, pages=page_count, chars=len(text), reviewed=False))
(OUT / 'extraction.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2)+'\n')
print(json.dumps(rows, ensure_ascii=False, indent=2))
