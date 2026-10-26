"""Build a reviewed local intake package; never authenticates or publishes."""
from pathlib import Path
import hashlib
import json
import shutil
from docx import Document

ROOT = Path(__file__).resolve().parents[2]
ARCHIVE = ROOT / '.local/research/service-lab-2026-10-10'
OUT = ROOT / 'docs/research'
DATE = '2026-10-10'
FILES = [
    ('普通全日制本、专科毕业生办理毕业证明书程序及材料.docx', 'https://jwc.tsguas.edu.cn/info/2034/5188.htm', 2),
    ('学生学籍数据修改申请表.docx', 'https://jwc.tsguas.edu.cn/info/2034/5189.htm', 1),
]

def digest(value):
    return hashlib.sha256(value).hexdigest()

def write(name, value):
    (OUT / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')

def build():
    ARCHIVE.mkdir(parents=True, exist_ok=True)
    files = []
    for name, source, pages in FILES:
        supplied = Path('/Users/zhangweiqi/Downloads') / name
        archived = ARCHIVE / name
        if archived.exists() and digest(archived.read_bytes()) != digest(supplied.read_bytes()):
            raise ValueError('Different copy already archived; do not overwrite')
        if not archived.exists():
            shutil.copy2(supplied, archived)
        doc = Document(archived)
        text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        tables = [[list(dict.fromkeys(c.text for c in r.cells)) for r in t.rows] for t in doc.tables]
        files.append(dict(name=name, sourceUrl=source, sha256=digest(archived.read_bytes()),
            archive=str(archived.relative_to(ROOT)), renderedPageCount=pages,
            textAndTablesReviewed=True, containsFilledPersonalData=False,
            captchaSolvedByAgent=False, renderedVisualPass=False,
            renderLimitation='Bundled renderer produced pages but CJK glyphs were missing; no layout pass claimed',
            paragraphText=text, tables=tables))

    def item(key, title, category, publisher, url, body, old=None, published=''):
        result = dict(key=key, title=title, category=category, publisher=publisher,
            sourceUrl=url, publishedAt=published, expiresAt='2026-11-10T00:00', content=body.strip())
        if old is not None:
            result['supersedesDocumentId'] = old
        result['sha256'] = digest(result['content'].encode())
        return result

    docs = [item('graduation-proof', '毕业证遗失损坏如何办理毕业证明书：材料、代办与办理时限',
        '毕业生办事', '教务处（用户提供公开Word原件）', FILES[0][1], '''
普通全日制本专科毕业生办理毕业证明书
2026-10-10核读用户提供的教务官网Word原件，含办理说明及空白申请表。官网页面标2024-09-25，附件自身未标制度版本日期，不称2026新规。适用普通全日制本、专科毕业生学历证书遗失或损坏后的证明书申请，不等同于学位证明、补发原毕业证或继续教育统一流程。

办理材料与步骤
本人提出申请，填写《办理毕业证明书申请表》，交验本人身份证原件及复印件。
文件对2002年后的学历证明书（学信网可查学历），要求交验学信网学历证书电子注册备案表，另交本人1.5寸蓝底近期免冠照片1张及同版电子照片。
对2002年前的学历证明书（学信网不可查学历），要求交验个人档案中能证明学历类型、层次、就读时间、毕结业结论及学校相关证明人等材料，另交入学录取名册（学校综合档案室提取）。
原文写“2002年前/后”，没有单独解释2002当年归属；不要自行补成“2002年及以后”。2002当年或学信网查询情况不符时，先向教务处确认适用材料，不替申请人断定资格。

能否代办、多久办理
本人不能亲自到校，可书面授权他人代办，附委托人和被委托人双方签字的委托书，以及被委托人身份证原件、复印件；并非只带一张委托书即可省去上述本人材料。
文件写材料齐全且经学校核实无误后，学校在5个工作日予以办理或答复。不是自聊天咨询或首次交不完整材料起5天必定拿证，也不是5个自然日。

申请表怎么填写
表格列姓名、性别、联系方式、通讯地址、专业、身份证号、在校学习起止年月；办理人申请栏填写证书遗失或损毁时间、原证书编号，申请人签名及日期。承诺栏注明原证书作废及证明书注明原编号和遗失证明字样；原件同时适用于遗失、损坏，具体证明书文字由学校核验办理。
文件未给出统一收费、邮寄承诺、具体收件房间或网上申请入口，不能补造。先通过教务处正式渠道确认收件安排。
原文及下载入口：https://jwc.tsguas.edu.cn/info/2034/5188.htm
本站隐私提示：身份证、学历备案表、照片、签名和原证书编号只交学校指定渠道，不上传公开问答或校园墙。这里仅提供空白材料说明，不代签、代审或自动提交。
核验日期2026-10-10；本站复核期限不是学校制度失效日。
'''), item('status-change', '学生学籍数据修改申请表：原数据新数据、证明材料与签章',
        '教务表格与办事', '教务处（用户提供公开空白Word）', FILES[1][1], '''
学生学籍数据修改申请表填写指引
2026-10-10核读用户从教务处下载的空白Word原件，未含任何个人填写内容。来源页标2024-09-25，附件本身未标版本日期；不是2026新发布制度。

填写什么
基本栏位：申请人姓名、曾用名、性别、身份证号、学号、层次、所在学院、专业、班级、联系电话。之后分别填写修改项、原数据、新数据，把需要更正的字段和前后值写清楚；不确定的栏位先问学院，不虚构信息。
申请人声明信息属实并签名、填日期。所在学院栏写审核意见并签字盖章、填日期；教务处（招生办公室）栏盖章、填日期。签章应由相应人员和部门完成，不能自行代签或把填完表当作修改成功。

需要什么证明
表下注要求附与申请事项相符的合法性证明材料。空白表没有逐项列举姓名、身份证号等更正各需哪些证明，也没有规定统一份数、费用、审批天数或在线办理入口，应向学院、教务处确认当前要求。不要把这张表当作成绩修改、转专业或学信网自主改数据的授权。
来源及下载：https://jwc.tsguas.edu.cn/info/2034/5189.htm
本站隐私提示：只查询空白栏位说明，无需在AI里提供真实身份证号、学号、电话、签名或证明扫描件；填写后的材料只交学校指定渠道。本网站不会替学生修改教务系统记录。
核验2026-10-10；本站复核日期不等于学校表格失效日期。
''')]

    prior = json.loads((OUT / 'jwc-delivered-pack-2026-10-10.json').read_text())['documents'][1]
    old_section = '本次只补齐上述3张申请表；毕业证明书等其他附件内容仍未在本批核验，不能凭目录臆造字段、材料和盖章流程。'
    assert old_section in prior['content']
    revised = prior['content'].replace(old_section,
        '三张休复退学表维持以下已核读指引。2026-10-10另收到毕业证明书办理说明及学籍数据修改空白Word，已分主题核读入库，不再属于未读附件。毕业证明书材料、代办及5个工作日办理或答复条件，请查对应专题及 https://jwc.tsguas.edu.cn/info/2034/5188.htm ；学籍数据更正的原数据、新数据、合法性证明和签章栏，请查 https://jwc.tsguas.edu.cn/info/2034/5189.htm 。不要混用这些不同事项的表格。')
    docs.append(item('academic-forms-index', prior['title'], prior['category'], prior['publisher'],
        prior['sourceUrl'], revised, old=85))
    docs.append(item('lab-student-safety', '基础实验实训：学生操作、课堂安全和事故报告要求',
        '实验室使用规则', '基础实验实训中心（2020版制度）',
        'https://jcsysxzx.tsguas.edu.cn/info/1039/1640.htm', '''
基础实验实训中心学生使用规则
所核网页分别发布于2020-10-23、2020-12-18、2020-11-20；2026-10-10复核。保留制度年份，不宣称已验证全校所有实验室现行规程。

进入与操作（2020-10-23安全制度）
使用前接受安全教育，学生在教师或技术人员指导下操作。危险实验须有防护和教师监护；精密大型设备由专人管理，未经批准不得使用或拆卸，也不得擅改防护装置。保持通道畅通，不私拉乱接电线或超负荷用电；隐患上报，整改后需检查合格再投入使用。
来源：https://jcsysxzx.tsguas.edu.cn/info/1039/1640.htm

课堂上找谁（2020-12-18责任制度）
任课教师是课堂安全第一责任人，课前检查和安全教育，课中巡查指导；学生安全员协助。课后记录实训情况、报告隐患、整理清点。不会使用时先问任课教师，不独自试错。
来源：https://jcsysxzx.tsguas.edu.cn/info/1039/1647.htm

发生意外（2020-11-20预案）
立即向教师、中心负责人等报告；严重伤害按学校应急安排联系120或校医院，后续由学校组织报告和调查。本段只说明校内联系流程，不提供带电检修、化学处置或医疗急救操作教程。
来源：https://jcsysxzx.tsguas.edu.cn/info/1039/1642.htm

本站边界：教学安全说明不发放准入资格，不替代安全培训考试、设备具体型号规程和教师现场指导。安全检查表附件未在本批取得，不凭标题编造检查项目。最新安排联系基础实验实训中心：022-28776330，jcsysxzx@tsguas.edu.cn（普通办公联系，不是已核实的24小时急救电话；来源https://jcsysxzx.tsguas.edu.cn/）。
'''))
    docs.append(item('lab-maintenance-learning', '实验设备保养谁负责、实验课指导书和答疑找谁',
        '实验教学与设备维护', '基础实验实训中心（2020及2022版）',
        'https://jcsysxzx.tsguas.edu.cn/info/1039/1641.htm', '''
实验设备保养与实验课学习支持
以下是基础实验实训中心已读网页的版本摘要，2026-10-10复核，不视为当前所有学院统一执行细则。

谁保养、多久一次（2022-03-18三级保养制度）
日常保养依说明书检查与清洁；一级保养每学期一次，由使用人员与维修人员共同完成；二级保养每学年一次，以专业维修人员为主。保养情况和问题应记录，不能把制度改写为学生可以自行拆机维修或操作带电设备。
来源：https://jcsysxzx.tsguas.edu.cn/info/1039/1641.htm

实验不会做、资料从哪里来（2020-10-15教学管理制度）
教师应提供必要参考资料并安排答疑。实验应有计划、教材或指导书，教学训练包括操作、测量处理和实验报告。学生可先向任课教师索取本课程指导书并请其指导疑难点。
来源：https://jcsysxzx.tsguas.edu.cn/info/1039/1646.htm

本站边界：本批未读设备保养记录单附件，不虚构字段或已完成保养。网页没有给出对外统一开放时段、自助预约入口或全校实验报告模板；“开放实验室”栏目所见参观活动不代表日常可自行入内。实际时间、准入和设备规程向任课教师、实验室负责人确认。中心办公电话022-28776330，邮箱jcsysxzx@tsguas.edu.cn（来源https://jcsysxzx.tsguas.edu.cn/）。
'''))
    docs.append(item('staff-forms-directory', '报销、学生酬金及教师办公表格官方入口',
        '办事表格目录', '基础实验实训中心（公开历史表单目录）',
        'https://jcsysxzx.tsguas.edu.cn/xzzx/cybg.htm', '''
教师及经办人员常用表格去哪里找
2026-10-10核对公开目录及页面附件名称，仅提供查找入口；附件内容未逐份核读，不将标题当作已经验证的填表流程，也不据旧模板确定2026报销资格、标准或审批权限。

财务：2024-10-11发布的30表目录，含差旅报销、支出报销审核、票据粘贴、学生酬金、财务档案借阅等。目录内还含2019、2020、2022等年份模板，发布于2024不等于全是2024新版。
入口：https://jcsysxzx.tsguas.edu.cn/info/1060/1590.htm
教师休假、销假与国内进修培训：2018年目录 https://jcsysxzx.tsguas.edu.cn/info/1060/1045.htm
档案借阅、印章使用、电话业务与会议申请等：2017年办公表格目录 https://jcsysxzx.tsguas.edu.cn/info/1060/1044.htm
总目录：https://jcsysxzx.tsguas.edu.cn/xzzx/cybg.htm

实际办理先向所属部门及财务、人事或相关归口部门核对当前模板与正式提交渠道。这里不提供代签、盖章或转账功能；学生酬金表是经办材料，不表示每个学生可直接据此申请或必然领取补助。填好的表可能含身份证、银行账号、金额与签名，只交学校授权渠道，不上传公共AI、校园墙或知识库。
'''))
    pack = dict(schemaVersion=1, checkedAt=DATE, requiresNormalAdminReview=True,
        allSiteComplete=False, documents=docs)
    write('service-lab-pack-2026-10-10.json', pack)
    evidence = dict(schemaVersion=1, checkedAt=DATE, files=files,
        allSiteLinksExhausted=False, wholeSiteDenominatorKnown=False,
        pages=[dict(url=f['sourceUrl'], status='attachment_reviewed_from_user_copy',
            evidence='public_docx_reviewed_sha256:'+f['sha256'], textWindowComplete=True,
            note='User supplied public blank original; original captcha history retained') for f in files])
    write('service-lab-docx-evidence-2026-10-10.json', evidence)
    return pack

if __name__ == '__main__':
    result = build()
    print(json.dumps({'documents':len(result['documents']), 'revisions':1, 'newTopics':5}))
