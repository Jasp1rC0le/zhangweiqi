"""Read supplied public attachments without executing content or publishing it."""
import hashlib
import json
import shutil
import subprocess
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / '.local/research/delivered-2026-10-11'
FILES = [
    ('alumni-query', '天津中德应用技术大学校友档案查询申请表.doc', 15085966),
    ('external-teachers', '天津中德应用技术大学关于印发《天津中德应用技术大学外聘教师管理办法（修订）》的通知.PDF', 3927567),
    ('course-exams', '关于印发《课程考试主要教学环节规范标准》的通知（津中德教务〔2023〕12号）.pdf', 3927554),
    ('enterprise-internship', '关于印发《本科生企业实习教学环节规范标准》的通知（津中德教务〔2023〕15号）.pdf', 3927553),
    ('major-transfer', '关于开展2026年学生转专业工作的通知.rar', 3947269),
    ('project-change', '附件7：校级项目事项变更申请表.docx', 3930675),
    ('project-awards', '附件6：项目成果获得称号统计.doc', 3930672),
    ('project-results', '附件5：项目成果统计.doc', 3930674),
    ('deferred-projects', '附件4：历年暂缓通过项目申报结题情况汇总表.doc', 3930673),
    ('project-summary', '附件3：项目申报结题情况汇总表.doc', 3930670),
    ('project-application', '附件2：申报书+.doc', 3930668),
    ('project-evidence', '附件2：结题佐证材料.doc', 3930669),
    ('project-completion', '附件1：结题书.doc', 3930671),
    ('archive-retention', '附件：三合一表.xlsx', 3927319),
]
PARENTS = {
    'alumni-query': 'https://zhda.tsguas.edu.cn/info/1047/1386.htm',
    'external-teachers': 'https://jwc.tsguas.edu.cn/info/1887/5006.htm',
    'course-exams': 'https://jwc.tsguas.edu.cn/info/1887/4994.htm',
    'enterprise-internship': 'https://jwc.tsguas.edu.cn/info/1887/4993.htm',
    'major-transfer': 'https://jwc.tsguas.edu.cn/info/1887/5264.htm',
    'archive-retention': 'https://zhda.tsguas.edu.cn/info/1051/1383.htm',
    'project-application': 'https://jwc.tsguas.edu.cn/info/2034/5181.htm',
}
OWNERS = {'jwc': 1178850125, 'zhda': 1182472388}
MISSING = [
    ('收费公示（2026.4.29）.xlsx', 'cwsj', 1178325157, 3946860, 'info/1098/1766.htm'),
    ('2025届本科毕业设计（论文）答辩资料.zip', 'jwc', 1178850125, 3940108, 'info/1887/5237.htm'),
    ('2026在校生户口迁移通知.docx', 'aqbw', 1703757346, 15087394, 'info/1004/1719.htm'),
    ('中德安全小程序在校生落户功能使用手册.docx', 'aqbw', 1703757346, 15087395, 'info/1004/1719.htm'),
    ('各学院负责学生户籍辅导员信息表.xlsx', 'aqbw', 1703757346, 15087396, 'info/1004/1719.htm'),
    ('新生户口常见问题及解答20260907.pdf', 'aqbw', 1703757346, 15087443, 'info/1008/1720.htm'),
    ('2026高职升本科专业课考试成绩查询通知.pdf', 'zs', 2036629888, 3945737, 'info/1008/1821.htm'),
    ('2026高职升本科专业课考试成绩复核申请表.pdf', 'zs', 2036629888, 3945738, 'info/1008/1821.htm'),
]


def download_url(host, owner, file_id):
    return (f'https://{host}.tsguas.edu.cn/system/_content/download.jsp'
            f'?owner={owner}&urltype=news.DownloadAttachUrl&wbfileid={file_id}')


def public_receipt(rows):
    entries = []
    for row in rows:
        parent = PARENTS.get(row['key'], 'https://jwc.tsguas.edu.cn/info/2034/5182.htm')
        host = parent.split('/')[2].split('.')[0]
        suffix = Path(row['archive']).suffix.lower()[1:]
        reviewed = row.get('humanReviewed') is True
        entries.append({'url': download_url(host, OWNERS[host], row['downloadId']),
                        'parentUrl': parent, 'title': row['file'], 'archive': row['archive'],
                        'status': 'attachment_reviewed_from_user_copy' if reviewed else 'attachment_extracted_from_user_copy',
                        'evidence': f'public_{suffix}_{"reviewed" if reviewed else "extracted"}_sha256:' + row['sha256'],
                        'textWindowComplete': reviewed,
                        'note': '原件核读与发布分别记录；重跑保留原有人工审核证据，不由提取推断上线'})
    # Earlier receipts recorded parent URLs only. Link their reviewed originals to
    # the download URLs so the same files are not requested a second time.
    earlier = json.loads((ROOT / 'docs/research/service-lab-docx-evidence-2026-10-10.json').read_text())
    prior_ids = { 'https://jwc.tsguas.edu.cn/info/2034/5188.htm': 3930681,
                  'https://jwc.tsguas.edu.cn/info/2034/5189.htm': 3930682 }
    for item in earlier['files']:
        if item['sourceUrl'] not in prior_ids or item.get('textAndTablesReviewed') is not True:
            continue
        archive_path = ROOT / item['archive']
        if hashlib.sha256(archive_path.read_bytes()).hexdigest() != item['sha256']:
            raise ValueError('Earlier original hash mismatch; do not clear pending item')
        entries.append({'url': download_url('jwc', OWNERS['jwc'], prior_ids[item['sourceUrl']]),
                        'parentUrl': item['sourceUrl'], 'title': item['name'],
                        'status': 'attachment_reviewed_from_user_copy',
                        'evidence': 'public_docx_reviewed_sha256:' + item['sha256'],
                        'textWindowComplete': True,
                        'note': '复用10月10日全文核读证据，仅补下载地址映射；不是重新导入'})
    additional_path = ROOT / 'docs/research/additional-attachments-received-2026-10-11.json'
    additional = json.loads(additional_path.read_text()) if additional_path.exists() else {}
    received = {download_key(item['downloadUrl']): item for item in additional.get('files', [])}
    unresolved = 0
    for title, host, owner, file_id, parent in MISSING:
        url = download_url(host, owner, file_id)
        item = received.get(download_key(url))
        if item and item.get('received') and item.get('extracted'):
            original = ROOT / item['archive']
            if hashlib.sha256(original.read_bytes()).hexdigest() != item['sha256']:
                raise ValueError('Received original hash mismatch; preserve pending item')
            reviewed = item.get('humanReviewed') is True
            suffix = original.suffix.lower()[1:]
            entries.append({'url': url, 'parentUrl': item['sourceUrl'], 'title': title,
                            'archive': item['archive'],
                            'status': 'attachment_reviewed_from_user_copy' if reviewed else 'attachment_extracted_from_user_copy',
                            'evidence': f'public_{suffix}_{"reviewed" if reviewed else "extracted"}_sha256:' + item['sha256'],
                            'textWindowComplete': reviewed,
                            'note': '用户已补齐原件；仅解决资料内容缺口，不表示远程验证码已解除或全站收齐'})
            continue
        unresolved += 1
        entries.append({'url': url,
                        'parentUrl': f'https://{host}.tsguas.edu.cn/{parent}', 'title': title,
                        'status': 'captcha_required', 'textWindowComplete': False,
                        'note': '公开原文下载入口已核对，返回验证码页面；尚未取得文件，不完成验证码'})
    return {'schemaVersion': 1, 'checkedAt': '2026-10-11',
            'allSiteLinksExhausted': False, 'wholeSiteDenominatorKnown': False,
            'pages': entries,
            'notes': ['14份收到且提取；核读和发布另记，不能根据哈希推断已经上线',
                      f'原8个验证码入口仍待用户提供原件{unresolved}个；不是全站分母',
                      '2027图像采集原文未显示附件下载入口，未编造文件ID或下载链接']}


def download_key(url):
    # School download parameters can appear in either query order.
    from urllib.parse import parse_qs, urlsplit
    parsed = urlsplit(url)
    query = parse_qs(parsed.query)
    return parsed.hostname, query.get('owner', [None])[0], query.get('wbfileid', [None])[0]


def retain_review(fresh, previous):
    """Only retain human evidence when the exact archived bytes still match."""
    if previous.get('sha256') != fresh.get('sha256'):
        return fresh
    result = {**previous, **fresh}
    fields = ('humanReviewed', 'published', 'reviewedAt', 'documentId',
              'publishedInSummaryDocumentId', 'publicationEvidence',
              'renderVerified', 'renderNote', 'privacyFlags')
    for field in fields:
        if field in previous:
            result[field] = previous[field]
    if 'members' in fresh:
        old_members = {item['sha256']: item for item in previous.get('members', [])}
        result['members'] = [retain_review(item, old_members.get(item['sha256'], {}))
                             for item in fresh['members']]
    return result


def archive(source, target):
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest() != digest:
        raise ValueError('Different archive exists; preserve both')
    if not target.exists():
        shutil.copy2(source, target)
    return digest


def extract(path):
    suffix = path.suffix.lower()
    pages, sheets = None, None
    if suffix == '.pdf':
        pdf = PdfReader(path)
        pages = len(pdf.pages)
        text = '\n'.join(f'=== PDF PAGE {i + 1} ===\n{p.extract_text() or ""}' for i, p in enumerate(pdf.pages))
    elif suffix == '.docx':
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        for i, table in enumerate(doc.tables):
            text += f'\n=== TABLE {i + 1} ===\n'
            text += '\n'.join(' | '.join(dict.fromkeys(c.text for c in row.cells)) for row in table.rows)
    elif suffix == '.doc':
        result = subprocess.run(['/usr/bin/textutil', '-convert', 'txt', '-stdout', str(path)],
                                capture_output=True, check=True, timeout=30)
        text = result.stdout.decode('utf-8')
    elif suffix == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        parts, sheets = [], []
        for sheet in book:
            if sheet.max_row > 5000 or sheet.max_column > 100:
                raise ValueError('Workbook exceeds review limit')
            lines = []
            for row in sheet:
                values = [f'{cell.coordinate}={cell.value}' for cell in row if cell.value is not None]
                if values:
                    lines.append(' | '.join(values))
            parts.append(f'=== SHEET {sheet.title} (state={sheet.sheet_state}) ===\n' + '\n'.join(lines))
            sheets.append({'name': sheet.title, 'rows': sheet.max_row, 'columns': sheet.max_column,
                           'nonemptyRows': len(lines), 'state': sheet.sheet_state})
        book.close()
        text = '\n'.join(parts)
    else:
        raise ValueError('Unsupported attachment')
    if not text.strip() or len(text) > 200000:
        raise ValueError('No bounded readable content')
    path.with_suffix('.txt').write_text(text)
    return {'pages': pages, 'sheets': sheets, 'chars': len(text),
            'textArchive': str(path.with_suffix('.txt').relative_to(ROOT)),
            'extracted': True, 'humanReviewed': False, 'published': False}


def unpack_rar(path):
    # The user-provided archive contains only seven named office documents.
    expected = {
        '附件3：2026年学生转专业拟接收名额一览表.xlsx',
        '附件4：转专业申请审批表.docx', '附件5：学生申请转出情况汇总表.xlsx',
        '附件6：拟接收学生情况汇总表.xlsx', '关于开展2026年学生转专业工作的通知.docx',
        '附件1：转专业要求.docx', '附件2：本科招生专业选考科目要求.xlsx',
    }
    listing = subprocess.run(['/usr/bin/bsdtar', '-tf', str(path)],
                             capture_output=True, check=True, text=True, timeout=20).stdout.splitlines()
    if len(listing) != len(set(listing)) or set(listing) != expected:
        raise ValueError('Unexpected archive members; stop for inspection')
    detail = subprocess.run(['/usr/bin/bsdtar', '-tvf', str(path)],
                            capture_output=True, check=True, text=True, timeout=20).stdout
    if any(not line.startswith('-') for line in detail.splitlines()):
        raise ValueError('Archive contains a non-regular file')
    # bsdtar's listing exposes member sizes before any extraction.
    sizes = [int(line.split()[4]) for line in detail.splitlines()]
    if len(sizes) != len(listing) or max(sizes) > 10 * 1024 * 1024 or sum(sizes) > 40 * 1024 * 1024:
        raise ValueError('Archive exceeds safe expanded size')
    directory = OUT / 'major-transfer-members'
    directory.mkdir(exist_ok=True)
    rows = []
    for name in listing:
        target = directory / name
        if not target.exists():
            with target.open('xb') as output:
                subprocess.run(['/usr/bin/bsdtar', '-xOf', str(path), name],
                               stdout=output, stderr=subprocess.PIPE, check=True, timeout=30)
        rows.append({'file': name, 'archive': str(target.relative_to(ROOT)),
                     'sha256': hashlib.sha256(target.read_bytes()).hexdigest(), **extract(target)})
    return rows


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    extraction_path = OUT / 'extraction.json'
    previous = json.loads(extraction_path.read_text()) if extraction_path.exists() else {}
    previous_rows = {row['sha256']: row for row in previous.get('files', [])}
    rows = []
    for key, name, file_id in FILES:
        source = Path('/Users/zhangweiqi/Downloads') / name
        target = OUT / (key + source.suffix.lower())
        digest = archive(source, target)
        row = {'key': key, 'file': name, 'downloadId': file_id,
               'archive': str(target.relative_to(ROOT)), 'sha256': digest}
        if target.suffix == '.rar':
            row.update(members=unpack_rar(target), extracted=True, humanReviewed=False, published=False)
        else:
            row.update(extract(target))
        rows.append(retain_review(row, previous_rows.get(digest, {})))
    result = {**previous, 'schemaVersion': 1, 'receivedOn': '2026-10-11', 'files': rows,
              'notes': ['原件只读，不执行宏或外部链接；提取不等于核读、审核发布或实际下载成功',
                        'RAR为通知正文加6项附件；原件收到14个，展开后的可读文件20个']}
    extraction_path.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    evidence = ROOT / 'docs/research/attachment-intake-evidence-2026-10-11.json'
    evidence.write_text(json.dumps(public_receipt(rows), ensure_ascii=False, indent=2) + '\n')
    print(json.dumps([{'key': r['key'], 'chars': r.get('chars'), 'pages': r.get('pages'),
                       'members': len(r.get('members', [])), 'sheets': r.get('sheets')} for r in rows],
                     ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
