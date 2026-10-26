#!/usr/bin/env python3
"""Archive exactly the supplied admissions files locally; no upload or public export.

Run with the original attachment directory. Originals are copied byte-for-byte,
not edited; SHA256 and complete extracted content remain outside public builds.
"""
import argparse
import hashlib
import json
import os
import shutil
from pathlib import Path

from docx import Document
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
PACK = ROOT / 'docs/research/admissions-2026.json'
ARCHIVE = ROOT / '.local/research/admissions-2026'


def archive(source_dir):
    pack = json.loads(PACK.read_text(encoding='utf-8'))
    inputs = [s for s in pack['sources'] if 'filename' in s]
    # Validate all inputs before making the archive, and refuse replacing a different file.
    for source in inputs:
        original = source_dir / source['filename']
        if not original.is_file():
            raise FileNotFoundError(original)
        target = ARCHIVE / source['archiveName']
        if target.exists() and original.read_bytes() != target.read_bytes():
            raise ValueError(f'Refusing to overwrite a different archive: {target.name}')
    ARCHIVE.mkdir(parents=True, exist_ok=True, mode=0o700)
    ARCHIVE.chmod(0o700)
    manifest = []
    for source in inputs:
        original = source_dir / source['filename']
        target = ARCHIVE / source['archiveName']
        if not target.exists():
            shutil.copyfile(original, target)
        target.chmod(0o600)
        digest = hashlib.sha256(target.read_bytes()).hexdigest()
        if target.suffix == '.pdf':
            units = [{'page': i + 1, 'text': p.extract_text() or ''}
                     for i, p in enumerate(PdfReader(target).pages)]
        else:
            doc = Document(target)
            units = [{'paragraph': i + 1, 'text': p.text}
                     for i, p in enumerate(doc.paragraphs) if p.text]
            units += [{'table': i + 1, 'rows': [[c.text for c in row.cells] for row in table.rows]}
                      for i, table in enumerate(doc.tables)]
        extracted = ARCHIVE / f"{source['id']}-extracted.json"
        extracted.write_text(json.dumps({'sourceId': source['id'], 'access': source['access'],
                                         'sha256': digest, 'units': units}, ensure_ascii=False, indent=2), encoding='utf-8')
        extracted.chmod(0o600)
        manifest.append({'id': source['id'], 'filename': source['filename'],
                         'archiveName': target.name, 'sha256': digest, 'bytes': target.stat().st_size,
                         'extractedFile': extracted.name, 'access': source['access'], 'units': len(units)})
    path = ARCHIVE / 'manifest.json'
    path.write_text(json.dumps({'receivedAt': pack['receivedAt'], 'sources': manifest,
                               'publicExport': False}, ensure_ascii=False, indent=2), encoding='utf-8')
    path.chmod(0o600)
    return {'archive': str(ARCHIVE), 'files': len(manifest), 'publicExport': False}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source_dir', type=Path)
    args = parser.parse_args()
    os.umask(0o077)
    print(json.dumps(archive(args.source_dir), ensure_ascii=False, indent=2))
